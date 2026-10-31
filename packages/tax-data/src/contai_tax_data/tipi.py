"""Contrato inicial para a tabela oficial TIPI em XLSX."""

from __future__ import annotations

import hashlib
import io
import json
import re
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

TIPI_SOURCE_URL = (
    "https://www.gov.br/receitafederal/pt-br/acesso-a-informacao/legislacao/"
    "documentos-e-arquivos/tipi.xlsx/%40%40download/file"
)
TIPI_ACCEPT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet, application/octet-stream"
TIPI_CONTRACT_VERSION = "tipi-v1"
_NCM_PATTERN = re.compile(r"^\d{4}\.\d{2}\.\d{2}$")


class TipiContractError(ValueError):
    """Entrada XLSX incompatível com o contrato TIPI suportado."""


class TipiQualityError(ValueError):
    """Dados TIPI fora dos limites de qualidade esperados."""


@dataclass(frozen=True, slots=True)
class TipiEntry:
    ncm_code: str
    ex_code: str | None
    description: str
    ipi_rate: str | None


@dataclass(frozen=True, slots=True)
class TipiDataset:
    entries: tuple[TipiEntry, ...]
    source_act: str


def _text(value: object) -> str:
    return " ".join(str(value or "").split())


def _rate(value: object) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str) and value.strip().upper() == "NT":
        return "NT"
    try:
        amount = Decimal(str(value).strip())
    except (InvalidOperation, ValueError):
        raise TipiContractError(f"alíquota inválida: {value!r}") from None
    if not amount.is_finite() or amount < 0:
        raise TipiContractError(f"alíquota inválida: {value!r}")
    return format(amount.normalize(), "f")


def parse_tipi_xlsx(body: bytes) -> TipiDataset:
    if not body:
        raise TipiContractError("arquivo XLSX vazio")
    try:
        workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as error:
        raise TipiContractError(f"arquivo XLSX inválido: {error}") from None
    try:
        sheet = next(
            (
                worksheet
                for worksheet in workbook.worksheets
                if any(
                    tuple(_text(cell).upper() for cell in row[:4]) == ("NCM", "EX", "DESCRIÇÃO", "ALÍQUOTA (%)")
                    for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, 40), values_only=True)
                )
            ),
            None,
        )
        if sheet is None:
            raise TipiContractError("cabeçalho NCM/EX/DESCRIÇÃO/ALÍQUOTA (%) não encontrado")

        entries: list[TipiEntry] = []
        seen: set[tuple[str, str | None]] = set()
        source_act_lines: list[str] = []
        source_act = ""
        header_seen = False
        for row in sheet.iter_rows(min_row=1, values_only=True):
            first = _text(row[0] if row else None)
            header = tuple(_text(cell).upper() for cell in row[:4]) if row else ()
            if not header_seen:
                if header == ("NCM", "EX", "DESCRIÇÃO", "ALÍQUOTA (%)"):
                    header_seen = True
                    source_act = "\n\n".join(source_act_lines)
                    continue
                if first:
                    source_act_lines.append(first)
                continue
            if not row or not isinstance(row[0], str):
                continue
            code = first
            if not _NCM_PATTERN.fullmatch(code):
                continue
            if len(row) < 4:
                raise TipiContractError(f"linha TIPI truncada para NCM {code}")
            ex_value = row[1]
            if ex_value is None or _text(ex_value) == "":
                ex_code = None
            elif isinstance(ex_value, int) and ex_value >= 0:
                ex_code = str(ex_value)
            elif isinstance(ex_value, str):
                ex_text = ex_value.strip()
                match = re.fullmatch(r"(?i:ex)\s*(\d+)", ex_text)
                if ex_text.isdigit():
                    ex_code = ex_text
                elif match:
                    ex_code = match.group(1)
                else:
                    raise TipiContractError(f"EX inválido para NCM {code}: {ex_value!r}")
            else:
                raise TipiContractError(f"EX inválido para NCM {code}: {ex_value!r}")
            description = _text(row[2])
            if not description:
                raise TipiContractError(f"descrição vazia para NCM {code}")
            key = (code, ex_code)
            if key in seen:
                raise TipiContractError(f"NCM/EX duplicado: {code}/{ex_code or 'sem EX'}")
            seen.add(key)
            entries.append(TipiEntry(code, ex_code, description, _rate(row[3])))
    finally:
        workbook.close()

    if not entries:
        raise TipiQualityError("nenhuma linha de produto TIPI encontrada")
    entries.sort(key=lambda item: (item.ncm_code, item.ex_code or ""))
    return TipiDataset(tuple(entries), source_act)


def validate_tipi(
    dataset: TipiDataset,
    *,
    minimum_entries: int = 8_000,
    maximum_entries: int = 14_000,
    minimum_ex_entries: int = 500,
    maximum_missing_rates: int = 10,
) -> None:
    count = len(dataset.entries)
    if not minimum_entries <= count <= maximum_entries:
        raise TipiQualityError(f"volume TIPI inesperado: {count} (esperado {minimum_entries}-{maximum_entries})")
    if sum(entry.ex_code is not None for entry in dataset.entries) < minimum_ex_entries:
        raise TipiQualityError("volume de linhas EX inesperado")
    if sum(entry.ipi_rate is None for entry in dataset.entries) > maximum_missing_rates:
        raise TipiQualityError("volume de alíquotas ausentes inesperado")


def tipi_snapshot(dataset: TipiDataset) -> bytes:
    lines = [
        json.dumps(asdict(entry), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        for entry in dataset.entries
    ]
    return ("\n".join(lines) + "\n").encode("utf-8")


def tipi_release_id(dataset: TipiDataset) -> str:
    identity = json.dumps(
        {"snapshot_sha256": hashlib.sha256(tipi_snapshot(dataset)).hexdigest(), "source_act": dataset.source_act},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(identity).hexdigest()
